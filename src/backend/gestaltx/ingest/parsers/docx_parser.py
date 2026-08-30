"""DOCX extraction using python-docx."""

from __future__ import annotations

from pathlib import Path

from ..metadata import infer_metadata
from ..models import Document, Section


def parse_docx(path: str | Path, corpus_root: str | Path | None = None) -> Document:
    """Extract headings, paragraphs, and simple key/value tables from a DOCX."""
    try:
        from docx import Document as DocxDocument
    except ImportError as exc:  # pragma: no cover - depends on optional package
        raise RuntimeError("DOCX parsing requires the 'python-docx' package") from exc

    source = Path(path)
    metadata = infer_metadata(source, corpus_root)
    raw = DocxDocument(str(source))
    title = metadata["title"]
    current_title = title
    body: list[str] = []
    section_data: list[tuple[str, str]] = []

    def flush() -> None:
        nonlocal body
        text = "\n\n".join(body).strip()
        if text:
            section_data.append((current_title, text))
        body = []

    for paragraph in raw.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name if paragraph.style else "").lower()
        if style.startswith("title") and not section_data and not body:
            title = text
            current_title = text
        elif style.startswith("heading"):
            flush()
            current_title = text
        else:
            body.append(text)
    flush()

    infobox: dict[str, str] = {}
    table_text: list[str] = []
    for table in raw.tables:
        for row in table.rows:
            cells = [" ".join(cell.text.split()) for cell in row.cells]
            if len(cells) == 2 and cells[0] and cells[1]:
                infobox.setdefault(cells[0].rstrip(":"), cells[1])
            if any(cells):
                table_text.append(" | ".join(cells))
    if table_text:
        section_data.append(("Tables", "\n".join(table_text)))

    if not section_data:
        section_data = [(title, "")]
    sections = [
        Section(
            section_id=f"{metadata['doc_id']}:section:{order:04d}",
            title=heading,
            text=text,
            order=order,
        )
        for order, (heading, text) in enumerate(section_data)
    ]
    subject = metadata["subject_entity"]
    metadata["title"] = title
    return Document(
        **metadata,
        sections=sections,
        entities=[subject] if subject else [],
        infobox=infobox,
        metadata={"extension": ".docx"},
    )


parse = parse_docx
